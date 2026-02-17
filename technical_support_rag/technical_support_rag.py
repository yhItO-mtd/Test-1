# technical_support_rag.py
"""
Technical Support RAG System - Streamlit Application

Documents (PDF, DOCX, PPTX, TXT) are uploaded, indexed with LlamaIndex,
and queried via a local Ollama LLM. All processing runs offline.
"""
import shutil
import json
from datetime import datetime
from pathlib import Path

import streamlit as st
import PyPDF2
import docx
from pptx import Presentation
from llama_index.core import (
    VectorStoreIndex,
    Document,
    Settings,
    StorageContext,
    load_index_from_storage,
)
from llama_index.embeddings.huggingface import HuggingFaceEmbedding
from llama_index.llms.ollama import Ollama

# ---------------------------------------------------------------------------
# Page config
# ---------------------------------------------------------------------------
st.set_page_config(
    page_title="Technical Support AI",
    page_icon="🔬",
    layout="wide",
    initial_sidebar_state="expanded",
)

STORAGE_DIR = Path("./storage")
METADATA_FILE = STORAGE_DIR / "metadata.json"

# ---------------------------------------------------------------------------
# Session state
# ---------------------------------------------------------------------------
if "index" not in st.session_state:
    st.session_state.index = None
if "chat_history" not in st.session_state:
    st.session_state.chat_history = []
if "documents" not in st.session_state:
    st.session_state.documents = []
# Cache extracted document parts so we never lose them across Streamlit reruns
if "doc_parts_cache" not in st.session_state:
    st.session_state.doc_parts_cache = {}  # filename -> list of part dicts

# ---------------------------------------------------------------------------
# Model setup
# ---------------------------------------------------------------------------

@st.cache_resource
def setup_models():
    """Initialise embedding model and LLM (cached across reruns)."""
    Settings.embed_model = HuggingFaceEmbedding(
        model_name="intfloat/multilingual-e5-large",
        cache_folder="./models",
    )
    Settings.llm = Ollama(
        model="llama3.1:8b",
        request_timeout=300.0,
        temperature=0.0,
        system_prompt=(
            "あなたは製品の技術サポート専門AIです。\n"
            "以下のルールに従ってください：\n"
            "1. 提供された文書の情報のみを使用する\n"
            "2. 不明な場合は「文書に記載がありません」と回答\n"
            "3. 具体的なページ番号や章を引用する\n"
            "4. 技術用語は正確に使用する\n"
            "5. 簡潔で分かりやすく回答する"
        ),
    )
    return True


try:
    setup_models()
    _models_ok = True
except Exception as exc:
    _models_ok = False
    st.error(f"モデル初期化エラー: {exc}")

# ---------------------------------------------------------------------------
# File extraction helpers
# ---------------------------------------------------------------------------

def extract_text_from_pdf(file) -> list[dict]:
    """Extract per-page text from a PDF."""
    pdf_reader = PyPDF2.PdfReader(file)
    documents = []
    for page_num, page in enumerate(pdf_reader.pages, 1):
        text = page.extract_text()
        if text and text.strip():
            documents.append(
                {
                    "text": text,
                    "metadata": {
                        "page": page_num,
                        "total_pages": len(pdf_reader.pages),
                    },
                }
            )
    return documents


def extract_text_from_docx(file) -> list[dict]:
    """Extract text from a Word document."""
    doc = docx.Document(file)
    text = "\n".join(para.text for para in doc.paragraphs if para.text.strip())
    if not text.strip():
        return []
    return [{"text": text, "metadata": {}}]


def extract_text_from_pptx(file) -> list[dict]:
    """Extract per-slide text from a PowerPoint file."""
    prs = Presentation(file)
    documents = []
    for slide_num, slide in enumerate(prs.slides, 1):
        text_parts = [
            shape.text
            for shape in slide.shapes
            if hasattr(shape, "text") and shape.text.strip()
        ]
        if text_parts:
            documents.append(
                {
                    "text": "\n".join(text_parts),
                    "metadata": {
                        "slide": slide_num,
                        "total_slides": len(prs.slides),
                    },
                }
            )
    return documents


def load_document(uploaded_file) -> list[dict] | None:
    """Dispatch to the correct extractor based on file extension."""
    ext = Path(uploaded_file.name).suffix.lower()
    try:
        if ext == ".pdf":
            return extract_text_from_pdf(uploaded_file)
        elif ext == ".docx":
            return extract_text_from_docx(uploaded_file)
        elif ext == ".pptx":
            return extract_text_from_pptx(uploaded_file)
        elif ext == ".txt":
            text = uploaded_file.read().decode("utf-8")
            if not text.strip():
                return []
            return [{"text": text, "metadata": {}}]
        else:
            st.warning(f"未対応のファイル形式: {ext}")
            return None
    except Exception as e:
        st.error(f"ファイル読み込みエラー ({uploaded_file.name}): {e}")
        return None

# ---------------------------------------------------------------------------
# Index helpers
# ---------------------------------------------------------------------------

def _build_index_from_cache() -> VectorStoreIndex | None:
    """Build a LlamaIndex VectorStoreIndex from the session doc_parts_cache."""
    all_documents: list[Document] = []
    for filename, parts in st.session_state.doc_parts_cache.items():
        # Find file_type from documents metadata
        file_type = ""
        for d in st.session_state.documents:
            if d["name"] == filename:
                file_type = d["type"]
                break

        for part in parts:
            metadata = {
                "file_name": filename,
                "file_type": file_type,
                **part["metadata"],
            }
            all_documents.append(Document(text=part["text"], metadata=metadata))

    if not all_documents:
        return None

    index = VectorStoreIndex.from_documents(all_documents)

    # Persist
    STORAGE_DIR.mkdir(parents=True, exist_ok=True)
    index.storage_context.persist(persist_dir=str(STORAGE_DIR))
    with open(METADATA_FILE, "w", encoding="utf-8") as f:
        json.dump(st.session_state.documents, f, ensure_ascii=False, indent=2)

    return index

# ---------------------------------------------------------------------------
# Sidebar
# ---------------------------------------------------------------------------
with st.sidebar:
    st.header("ナレッジベース管理")

    # Load persisted index
    if STORAGE_DIR.exists() and st.session_state.index is None:
        if st.button("保存済みデータを読み込む"):
            with st.spinner("読み込み中..."):
                try:
                    storage_context = StorageContext.from_defaults(
                        persist_dir=str(STORAGE_DIR)
                    )
                    st.session_state.index = load_index_from_storage(storage_context)
                    if METADATA_FILE.exists():
                        with open(METADATA_FILE, "r", encoding="utf-8") as f:
                            st.session_state.documents = json.load(f)
                    st.success("読み込み完了")
                except Exception as exc:
                    st.error(f"読み込みエラー: {exc}")

    st.divider()

    # File uploader
    uploaded_files = st.file_uploader(
        "文書をアップロード",
        type=["pdf", "docx", "pptx", "txt"],
        accept_multiple_files=True,
        help="取扱説明書、セミナー資料、アプリケーションノート等",
    )

    if uploaded_files:
        new_docs_added = False

        for uploaded_file in uploaded_files:
            fname = uploaded_file.name
            if fname in st.session_state.doc_parts_cache:
                # Already processed
                continue

            # Reset stream position so the extractor reads from the start
            uploaded_file.seek(0)
            doc_parts = load_document(uploaded_file)

            if doc_parts:
                # Cache the extracted parts in session state so they survive reruns
                st.session_state.doc_parts_cache[fname] = doc_parts
                st.session_state.documents.append(
                    {
                        "name": fname,
                        "type": Path(fname).suffix[1:].upper(),
                        "uploaded_at": datetime.now().isoformat(),
                        "parts": len(doc_parts),
                    }
                )
                new_docs_added = True

        if new_docs_added or st.session_state.index is None:
            with st.spinner("インデックス構築中..."):
                idx = _build_index_from_cache()
                if idx is not None:
                    st.session_state.index = idx
                    st.success(f"{len(st.session_state.documents)}件の文書を登録")

    # Registered documents
    if st.session_state.documents:
        st.subheader("登録済み文書")
        _icons = {"PDF": "📕", "DOCX": "📘", "PPTX": "📊", "TXT": "📄"}
        for doc in st.session_state.documents:
            icon = _icons.get(doc["type"], "📄")
            st.text(f"{icon} {doc['name']}")
            if "parts" in doc:
                st.caption(f"   {doc['parts']}セクション")

        st.divider()

        if st.button("すべてクリア", type="secondary"):
            st.session_state.documents = []
            st.session_state.index = None
            st.session_state.chat_history = []
            st.session_state.doc_parts_cache = {}
            if STORAGE_DIR.exists():
                shutil.rmtree(STORAGE_DIR)
            st.rerun()

# ---------------------------------------------------------------------------
# Main area
# ---------------------------------------------------------------------------
st.title("Technical Support AI")
st.caption("取扱説明書・セミナー資料から自動回答")

# Chat history
for message in st.session_state.chat_history:
    with st.chat_message(message["role"]):
        st.markdown(message["content"])

        if message.get("sources"):
            with st.expander("参照元の詳細", expanded=False):
                for i, source in enumerate(message["sources"], 1):
                    st.markdown(f"**[{i}] {source['file_name']}**")
                    if source.get("page"):
                        st.info(
                            f"ページ {source['page']}/{source.get('total_pages', '?')}"
                        )
                    elif source.get("slide"):
                        st.info(
                            f"スライド {source['slide']}/{source.get('total_slides', '?')}"
                        )
                    if source.get("score") is not None:
                        st.caption(f"関連度: {source['score']:.3f}")
                    st.markdown("**参照した文章（原文）：**")
                    st.text_area(
                        f"原文_{i}",
                        source["text"],
                        height=150,
                        key=f"source_{message.get('timestamp', '')}_{i}",
                        label_visibility="collapsed",
                    )
                    st.divider()

# Query input
if st.session_state.index is not None:
    st.markdown("### よくある質問")
    col1, col2, col3 = st.columns(3)
    with col1:
        if st.button("仕様・スペック"):
            st.session_state.quick_question = (
                "この製品の主な仕様とスペックを教えてください"
            )
    with col2:
        if st.button("トラブルシューティング"):
            st.session_state.quick_question = (
                "よくあるトラブルとその対処法を教えてください"
            )
    with col3:
        if st.button("メンテナンス"):
            st.session_state.quick_question = (
                "日常的なメンテナンス方法を教えてください"
            )

    prompt = st.chat_input("質問を入力してください（例：ベースライン補正の手順は？）")

    if "quick_question" in st.session_state:
        prompt = st.session_state.quick_question
        del st.session_state.quick_question

    if prompt:
        timestamp = datetime.now().isoformat()

        st.session_state.chat_history.append(
            {"role": "user", "content": prompt, "timestamp": timestamp}
        )
        with st.chat_message("user"):
            st.markdown(prompt)

        with st.chat_message("assistant"):
            with st.spinner("文書を検索中..."):
                try:
                    query_engine = st.session_state.index.as_query_engine(
                        similarity_top_k=5,
                        response_mode="compact",
                    )
                    response = query_engine.query(prompt)
                except Exception as exc:
                    st.error(
                        f"回答生成エラー: {exc}\n\n"
                        "Ollamaが起動しているか確認してください: `ollama serve`"
                    )
                    st.stop()

                st.markdown(response.response)

                sources: list[dict] = []
                for node in response.source_nodes:
                    sources.append(
                        {
                            "file_name": node.metadata.get("file_name", "unknown"),
                            "file_type": node.metadata.get("file_type", ""),
                            "page": node.metadata.get("page"),
                            "slide": node.metadata.get("slide"),
                            "total_pages": node.metadata.get("total_pages"),
                            "total_slides": node.metadata.get("total_slides"),
                            "score": node.score,
                            "text": node.text,
                        }
                    )

                if sources:
                    with st.expander("参照元の詳細", expanded=True):
                        for i, source in enumerate(sources, 1):
                            st.markdown(f"**[{i}] {source['file_name']}**")
                            if source.get("page"):
                                st.info(
                                    f"ページ {source['page']}/{source.get('total_pages', '?')}"
                                )
                            elif source.get("slide"):
                                st.info(
                                    f"スライド {source['slide']}/{source.get('total_slides', '?')}"
                                )
                            if source.get("score") is not None:
                                st.caption(f"関連度: {source['score']:.3f}")
                            st.markdown("**参照した文章（原文）：**")
                            st.text_area(
                                f"原文_{i}",
                                source["text"],
                                height=150,
                                key=f"source_new_{timestamp}_{i}",
                                label_visibility="collapsed",
                            )
                            st.divider()

                st.session_state.chat_history.append(
                    {
                        "role": "assistant",
                        "content": response.response,
                        "sources": sources,
                        "timestamp": timestamp,
                    }
                )
else:
    st.info("サイドバーから取扱説明書やセミナー資料をアップロードしてください")

    st.markdown(
        """
### 使い方

1. **文書をアップロード**: PDF、Word、PowerPoint形式の技術文書
2. **質問を入力**: 製品仕様、操作方法、トラブル対処など
3. **回答を確認**: 参照元の原文も確認できます

### 対応文書
- 📕 取扱説明書（PDF）
- 📊 セミナー資料（PPTX）
- 📘 アプリケーションノート（Word/PDF）
- 📄 技術資料（TXT）
"""
    )

# ---------------------------------------------------------------------------
# Footer
# ---------------------------------------------------------------------------
st.divider()
col1, col2, col3, col4 = st.columns(4)

with col1:
    if st.button("会話履歴クリア"):
        st.session_state.chat_history = []
        st.rerun()

with col2:
    if st.session_state.chat_history:
        export_data = {
            "export_date": datetime.now().isoformat(),
            "conversation": st.session_state.chat_history,
        }
        st.download_button(
            "会話をエクスポート",
            json.dumps(export_data, ensure_ascii=False, indent=2),
            f"conversation_{datetime.now().strftime('%Y%m%d_%H%M%S')}.json",
            "application/json",
        )

with col3:
    if st.session_state.index:
        st.success("準備完了")
    else:
        st.warning("文書待機中")

with col4:
    st.caption(f"登録: {len(st.session_state.documents)}件")
